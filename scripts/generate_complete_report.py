"""
generate_complete_report.py
Constructs the complete, exhaustive, publication-grade Final Examination Report for:
24CYS401 Secure Software Engineering Laboratory End Semester Examination.
System 23: Autonomous Warehouse Management System (AWMS)

Covers EVERY SINGLE sub-point across all 16 Phases:
- Phase 1: Agile Process & Development Approach [6 Marks]
- Phase 2: Requirements Engineering [7 Marks]
- Phase 3: Requirements Analysis & UML [7 Marks]
- Phase 4: Data & Information Flow Modeling [7 Marks]
- Phase 5: Software Architecture & Design Engineering [7 Marks]
- Phase 6: User Interface Design [5 Marks]
- Phase 7: Threat Modeling & Security Analysis [10 Marks]
- Phase 8: Attack Tree & Security Architecture Refinement [6 Marks] (Includes Refactoring Code)
- Phase 9: Product Backlog & Jira/Scrum [7 Marks]
- Phase 10: Sprint Execution & Scrum Metrics [7 Marks]
- Phase 11: Secure Development & Build Environment [6 Marks]
- Phase 12: Secure Coding & Refactoring [4 Marks]
- Phase 13: Containerized Development: Docker & Kubernetes [7 Marks]
- Phase 14: CI/CD & Security Testing [7 Marks]
- Phase 15: Logging, Monitoring, Hardening & Secure Deployment [5 Marks]
- Phase 16: Final Security Review [1 Mark]
- FINAL SECTION: Audited Inventory of Software Tools & Exact Versions
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GRAPHS_DIR = os.path.join(BASE_DIR, "graphs")
MD_PATH = os.path.join(BASE_DIR, "FINAL_FULL_EXAM_REPORT.md")
DOCX_PATH = os.path.join(BASE_DIR, "FINAL_FULL_EXAM_REPORT.docx")

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_heading_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    if level == 1:
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(5)
        for run in h.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(16, 35, 63)
    elif level == 2:
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(3)
        for run in h.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(27, 73, 101)
    elif level == 3:
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        for run in h.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(60, 60, 60)
    return h

def add_image_safe(doc, image_name, caption_text, width_inches=5.6):
    img_path = os.path.join(GRAPHS_DIR, image_name)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(1)
        p_cap.paragraph_format.space_after = Pt(8)
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = 'Calibri'
        run_cap.font.size = Pt(8.5)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(100, 100, 100)

def add_table_styled(doc, headers, rows_data):
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_row = tbl.rows[0]
    for idx, h_text in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.text = h_text
        set_cell_background(cell, "10233F")
        set_cell_margins(cell, top=90, bottom=90, left=70, right=70)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    for r_idx, r_val in enumerate(rows_data, start=1):
        row = tbl.rows[r_idx]
        bg = "FFFFFF" if r_idx % 2 != 0 else "F7FAFC"
        for c_idx, c_val in enumerate(r_val):
            cell = row.cells[c_idx]
            cell.text = str(c_val)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
            p = cell.paragraphs[0]
            for r in p.runs:
                r.font.size = Pt(7.5)
                r.font.name = 'Calibri'
                if c_idx == 0:
                    r.font.bold = True
                    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl

def add_code_block(doc, code_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(code_str)
    r.font.name = 'Consolas'
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor(20, 20, 20)

print("Starting generation...")
